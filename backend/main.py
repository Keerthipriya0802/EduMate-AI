import json
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import fitz
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pptx import Presentation
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
APP_DATA = ROOT / "app-data"
UPLOAD_DIR = APP_DATA / "uploads"
STATE_PATH = APP_DATA / "state.json"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
APP_DATA.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="EduMate AI API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def utc_now():
    return datetime.utcnow().isoformat(timespec="seconds") + "Z"


def tokenize(text: str):
    return re.findall(r"[a-zA-Z][a-zA-Z0-9-]{1,}", text.lower())


def load_state() -> dict[str, Any]:
    if not STATE_PATH.exists():
        return seed_demo_state()
    try:
        return json.loads(STATE_PATH.read_text())
    except json.JSONDecodeError:
        return seed_demo_state()


def save_state(state: dict[str, Any]):
    STATE_PATH.write_text(json.dumps(state, indent=2, ensure_ascii=False))


def make_chunk(text: str, page: int, doc_name: str, doc_type: str, subject: str, topic: str):
    words = tokenize(text)
    return {
        "id": f"chunk-{abs(hash(text + str(page)))}",
        "docName": doc_name,
        "docType": doc_type,
        "subject": subject,
        "topic": topic,
        "page": page,
        "text": text[:900],
        "keywords": list(dict.fromkeys(words))[:12],
    }


def extract_pdf_text(file_path: Path):
    doc = fitz.open(file_path)
    extracted = []
    for page_idx in range(doc.page_count):
        page = doc[page_idx]
        text = page.get_text("text")
        if text.strip():
            extracted.append({"page": page_idx + 1, "text": text.strip()})
    doc.close()
    return extracted


def extract_text_from_file(file_path: Path):
    suffix = file_path.suffix.lower()
    if suffix == ".txt":
        return [{"page": 1, "text": file_path.read_text(encoding="utf-8", errors="ignore")}]
    if suffix == ".pdf":
        return extract_pdf_text(file_path)
    if suffix == ".pptx":
        prs = Presentation(str(file_path))
        slides = []
        for idx, slide in enumerate(prs.slides, start=1):
            text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text)
            if text.strip():
                slides.append({"page": idx, "text": text.strip()})
        return slides
    raise ValueError("Unsupported file format for extraction. Supported formats: PDF, TXT, PPTX.")


def generate_demo_pdf(path: Path):
    pdf = canvas.Canvas(str(path), pagesize=(595, 842))
    pdf.setTitle("Database Systems - Normalization")
    pdf.setAuthor("EduMate AI")
    lines = [
        "Database Normalization",
        "Normalization is the process of organizing data to reduce redundancy and improve integrity.",
        "In a well-designed database, each fact should appear in one place. This makes updates easier and reduces anomalies.",
        "", 
        "1NF: Each table cell contains a single value and each row is unique.",
        "2NF: A table is in 2NF when it is already in 1NF and all non-key attributes depend on the full primary key.",
        "3NF: A table is in 3NF when it is in 2NF and all non-key attributes depend only on the primary key, not on other non-key attributes.",
        "",
        "Example: A StudentCourse table may store student_id, student_name, course_id, course_name, and instructor_name.",
        "This design repeats student and course details. By separating the data into Student, Course, and Enrollment tables, normalization removes duplication.",
        "A student who enrolls in two courses should not have to repeat the same student information in multiple rows.",
        "", 
        "Normalization improves consistency, reduces update anomalies, and keeps relational data easier to query.",
    ]
    y = 780
    for line in lines:
        if y < 60:
            pdf.showPage()
            y = 780
        pdf.drawString(56, y, line[:110])
        y -= 28
    pdf.save()


def seed_demo_state() -> dict[str, Any]:
    file_name = "dbms-textbook-normalization.pdf"
    file_path = UPLOAD_DIR / file_name
    if not file_path.exists():
        generate_demo_pdf(file_path)
    pages = extract_pdf_text(file_path)
    chunks = []
    for page_data in pages:
        chunks.append(make_chunk(page_data["text"], page_data["page"], file_name, "pdf", "Database Systems", "Normalization"))

    state = {
        "users": [
            {
                "id": 1,
                "name": "Aarav Sharma",
                "email": "aarav@example.com",
                "password": "password123",
                "studyGoalHours": 12,
                "course": "B.Tech Computer Science",
            }
        ],
        "materials": [
            {
                "id": "material-1",
                "name": file_name,
                "filePath": str(file_path),
                "subject": "Database Systems",
                "chapter": "Relational Design",
                "topic": "Normalization",
                "type": "pdf",
                "uploadedAt": utc_now(),
                "status": "processed",
                "uploadedByUserId": 1,
                "chunks": chunks,
            }
        ],
        "chatSessions": [{
            "id": "session-1",
            "userId": 1,
            "title": "Normalization revision",
            "messages": [],
        }],
        "quizHistory": [
            {
                "quizId": "quiz-1",
                "userId": 1,
                "topic": "Normalization",
                "score": 60,
                "submittedAt": (datetime.utcnow() - timedelta(days=3)).isoformat(timespec="seconds") + "Z",
            }
        ],
        "topicMastery": [
            {"userId": 1, "topic": "Normalization", "mastery": 72, "trend": 8},
            {"userId": 1, "topic": "SQL Queries", "mastery": 84, "trend": 5},
            {"userId": 1, "topic": "ER Modeling", "mastery": 67, "trend": -2},
            {"userId": 1, "topic": "Transactions", "mastery": 81, "trend": 4},
        ],
        "recommendations": [
            {
                "id": "rec-1",
                "userId": 1,
                "topic": "Normalization",
                "reason": "Two recent quiz questions on 2NF and 3NF were incorrect.",
                "duration": 28,
                "priority": "high",
                "actionLabel": "Start Learning",
            },
            {
                "id": "rec-2",
                "userId": 1,
                "topic": "ER Modeling",
                "reason": "Weakest topic in the last revision cycle.",
                "duration": 20,
                "priority": "medium",
                "actionLabel": "Review Notes",
            }
        ],
        "revisionTasks": [
            {"id": "task-1", "userId": 1, "title": "Revise 2NF examples", "topic": "Normalization", "dueDate": "2026-10-15", "status": "pending"},
            {"id": "task-2", "userId": 1, "title": "Practice SQL joins", "topic": "SQL Queries", "dueDate": "2026-10-16", "status": "done"},
            {"id": "task-3", "userId": 1, "title": "Summarize transaction states", "topic": "Transactions", "dueDate": "2026-10-17", "status": "pending"},
        ],
    }
    save_state(state)
    return state


def get_user_by_id(user_id: int):
    state = load_state()
    for user in state["users"]:
        if user["id"] == user_id:
            return user
    return None


def build_dashboard(user_id: int):
    state = load_state()
    user = get_user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
    topic_mastery = state["topicMastery"]
    mastery_avg = round(sum(item["mastery"] for item in topic_mastery if item["userId"] == user_id) / max(len([item for item in topic_mastery if item["userId"] == user_id]), 1), 1)
    progress = min(100, max(0, int(mastery_avg)))
    return {
        "welcome": f"Welcome back, {user['name'].split()[0]}!",
        "studyHours": 14.5,
        "completedLessons": 18,
        "overallProgress": progress,
        "dailyGoal": 12,
        "goalProgress": 78,
        "weakTopics": ["Normalization", "ER Modeling"],
        "currentCourses": ["Database Systems", "Operating Systems", "Algorithms"],
        "topicMastery": [{"topic": item["topic"], "value": item["mastery"]} for item in topic_mastery if item["userId"] == user_id],
        "recentMaterials": [
            {"name": material["name"], "subject": material["subject"], "updated": "2h ago"} for material in state["materials"][:3]
        ],
        "upcomingRevision": [
            {"name": task["title"], "time": task["dueDate"]} for task in state["revisionTasks"][:3]
        ],
    }


def get_materials_for_user(user_id: int):
    state = load_state()
    return [
        {
            "id": material["id"],
            "name": material["name"],
            "subject": material["subject"],
            "chapter": material["chapter"],
            "topic": material["topic"],
            "type": material["type"],
            "uploadedAt": material["uploadedAt"],
            "status": material["status"],
            "uploadedByUserId": material["uploadedByUserId"],
        }
        for material in state["materials"]
        if material["uploadedByUserId"] == user_id
    ]


def find_relevant_chunks(question: str, materials: list[dict[str, Any]], limit: int = 3):
    keywords = tokenize(question)
    ranked = []
    for material in materials:
        for chunk in material.get("chunks", []):
            chunk_text = (chunk.get("text") or "").lower()
            score = sum(3 for word in keywords if word in chunk_text)
            score += sum(2 for word in keywords if word in (material.get("topic") or "").lower())
            if score > 0:
                ranked.append({"score": score, "material": material, "chunk": chunk})
    ranked.sort(key=lambda item: item["score"], reverse=True)
    return ranked[:limit]


def build_answer(question: str, level: str = "beginner", materials: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    materials = materials if materials is not None else load_state()["materials"]
    matched = find_relevant_chunks(question, materials)
    if not matched:
        return {
            "answer": "I could not find enough reliable information in the uploaded study materials to answer this confidently. Please upload a related DBMS chapter or ask a question that matches the uploaded notes.",
            "citations": [],
            "mode": "demo",
            "followUps": [
                "What is 1NF in database normalization?",
                "Explain the difference between 2NF and 3NF."
            ],
        }

    primary = matched[0]
    material = primary["material"]
    chunk = primary["chunk"]
    text = chunk["text"]
    q = question.lower()
    if "normalization" in q:
        answer = (
            "Database normalization is the process of organizing table data so it has no unnecessary duplication and each fact is stored in one place. "
            "A simple example is a StudentCourse table that stores student_id, student_name, course_id, course_name, and instructor_name. "
            "That repeats student and course data across rows. By splitting it into separate Student, Course, and Enrollment tables, we reduce redundancy and avoid update anomalies."
        )
    elif "2nf" in q or "second normal form" in q:
        answer = "2NF means the table is already in 1NF and every non-key field depends on the entire primary key. This matters when a table has a composite key and one attribute depends only on part of it."
    elif "3nf" in q or "third normal form" in q:
        answer = "3NF means the table is in 2NF and no non-key attribute depends on another non-key attribute. In practice, we remove transitive dependencies so the data model stays cleaner and easier to maintain."
    else:
        answer = "The uploaded material explains the topic using a relational design example. The key idea is to separate repeated facts into smaller tables so each data item is stored once and relationships remain clear."

    if level == "advanced":
        answer += " From a design perspective, normalization reduces update anomalies, preserves referential integrity, and makes transactional consistency easier to enforce."
    elif level == "intermediate":
        answer += " This helps you avoid contradictions when the same data is updated in multiple places."
    else:
        answer += " Think of it as keeping your database organized like a clean spreadsheet where each fact lives in one logical place."

    citations = [{
        "documentName": material["name"],
        "page": chunk["page"],
        "sourceType": material["type"],
        "snippet": text[:180],
    }]

    return {
        "answer": answer,
        "citations": citations,
        "mode": "demo",
        "followUps": [
            "Can you give a real-world example of 1NF, 2NF, and 3NF?",
            "How does normalization reduce update anomalies?",
            "Test my understanding with a short quiz on normalization."
        ],
    }


def compute_quiz_questions(topic: str, difficulty: str, count: int):
    if topic.lower() == "normalization":
        base = [
            {
                "question": "Which normal form removes partial dependency in a relation?",
                "type": "multiple_choice",
                "options": ["1NF", "2NF", "3NF", "BCNF"],
                "answer": "2NF",
                "explanation": "2NF removes partial dependency by ensuring all non-key attributes depend on the complete primary key.",
            },
            {
                "question": "What is the main purpose of normalization?",
                "type": "multiple_choice",
                "options": ["Increase storage size", "Reduce redundancy and anomalies", "Add indexes automatically", "Duplicate rows in tables"],
                "answer": "Reduce redundancy and anomalies",
                "explanation": "Normalization organizes data to reduce duplication and improve integrity.",
            },
            {
                "question": "True or false: 3NF allows transitive dependencies between non-key attributes.",
                "type": "true_false",
                "options": ["True", "False"],
                "answer": "False",
                "explanation": "3NF prevents transitive dependencies, so non-key attributes should not depend on other non-key attributes.",
            },
            {
                "question": "A table in 1NF must ensure that each cell contains only one value and each row is unique. This is an example of:",
                "type": "multiple_choice",
                "options": ["Structural consistency", "Redundant data", "Index fallback", "Join explosion"],
                "answer": "Structural consistency",
                "explanation": "1NF focuses on atomic values and unique row identity.",
            },
        ]
    else:
        base = [
            {
                "question": f"Review the main concept for {topic}.",
                "type": "multiple_choice",
                "options": ["A", "B", "C", "D"],
                "answer": "A",
                "explanation": "This conceptual check ensures the topic is understood in context.",
            }
        ]
    if difficulty == "easy":
        selected = base[:2]
    elif difficulty == "hard":
        selected = base[:4]
    else:
        selected = base[:3]
    return selected[:count]


def update_mastery(user_id: int, topic: str, accuracy: float):
    state = load_state()
    found = False
    for item in state["topicMastery"]:
        if item["userId"] == user_id and item["topic"] == topic:
            item["mastery"] = max(0, min(100, round(item["mastery"] * 0.7 + accuracy * 100 * 0.3)))
            item["trend"] = max(-10, min(10, int((accuracy * 100) - 50)))
            found = True
            break
    if not found:
        state["topicMastery"].append({"userId": user_id, "topic": topic, "mastery": round(accuracy * 100), "trend": 5})
    save_state(state)


def build_recommendations(user_id: int):
    state = load_state()
    user_recs = [item for item in state["recommendations"] if item["userId"] == user_id]
    if user_recs:
        return user_recs
    weak_topic = sorted(
        [item for item in state["topicMastery"] if item["userId"] == user_id],
        key=lambda item: item["mastery"],
    )[0]
    return [{
        "id": "rec-auto",
        "userId": user_id,
        "topic": weak_topic["topic"],
        "reason": "Past assessment scores show this topic is consistently weak.",
        "duration": 25,
        "priority": "high",
        "actionLabel": "Start Learning",
    }]


def build_analytics(user_id: int):
    state = load_state()
    mastery = [item for item in state["topicMastery"] if item["userId"] == user_id]
    attempts = [item for item in state["quizHistory"] if item["userId"] == user_id]
    return {
        "mastery": {"labels": [item["topic"] for item in mastery], "values": [item["mastery"] for item in mastery]},
        "scoreTrend": [{"label": "Week 1", "score": 60}, {"label": "Week 2", "score": 65}, {"label": "Week 3", "score": 72}, {"label": "Week 4", "score": 78}, {"label": "Week 5", "score": 83}],
        "studyHours": [{"label": "Mon", "hours": 1.2}, {"label": "Tue", "hours": 2.5}, {"label": "Wed", "hours": 1.8}, {"label": "Thu", "hours": 3.1}, {"label": "Fri", "hours": 2.7}, {"label": "Sat", "hours": 2.2}],
        "completionRate": 78,
        "completedTasks": len([task for task in state["revisionTasks"] if task["userId"] == user_id and task["status"] == "done"]),
        "pendingTasks": len([task for task in state["revisionTasks"] if task["userId"] == user_id and task["status"] != "done"]),
        "quizAttempts": len(attempts),
    }


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "EduMate AI"}


@app.post("/api/auth/register")
def register_user(name: str = Form(...), email: str = Form(...), password: str = Form(...)):
    state = load_state()
    existing = any(user["email"] == email for user in state["users"])
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered.")
    user = {"id": len(state["users"]) + 1, "name": name, "email": email, "password": password, "studyGoalHours": 12, "course": "B.Tech Computer Science"}
    state["users"].append(user)
    save_state(state)
    return {"user": user}


@app.post("/api/auth/login")
def login_user(email: str = Form(...), password: str = Form(...)):
    state = load_state()
    user = next((item for item in state["users"] if item["email"] == email and item["password"] == password), None)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid credentials.")
    return {"user": user}


@app.get("/api/dashboard/{user_id}")
def dashboard(user_id: int):
    return build_dashboard(user_id)


@app.get("/api/materials/{user_id}")
def materials(user_id: int):
    return get_materials_for_user(user_id)


@app.post("/api/materials/upload")
def upload_material(file: UploadFile = File(...), subject: str = Form("Database Systems"), chapter: str = Form("General"), topic: str = Form("General"), user_id: int = Form(1)):
    if file.filename is None:
        raise HTTPException(status_code=400, detail="No file uploaded.")
    suffix = Path(file.filename).suffix.lower()
    if suffix not in {".pdf", ".txt", ".pptx"}:
        raise HTTPException(status_code=400, detail="Unsupported file type. Use PDF, TXT or PPTX files.")

    file_path = UPLOAD_DIR / file.filename
    content = file.file.read()
    file_path.write_bytes(content)
    try:
        extracted = extract_text_from_file(file_path)
    except (OSError, ValueError, RuntimeError) as exc:
        file_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=f"Could not process the uploaded file: {exc}") from exc

    state = load_state()
    material = {
        "id": f"material-{len(state['materials']) + 1}",
        "name": file.filename,
        "filePath": str(file_path),
        "subject": subject,
        "chapter": chapter,
        "topic": topic,
        "type": suffix.replace('.', '') or "txt",
        "uploadedAt": utc_now(),
        "status": "processed",
        "uploadedByUserId": user_id,
        "chunks": [
            make_chunk(page_data["text"], page_data["page"], file.filename, suffix.replace('.', '') or "txt", subject, topic)
            for page_data in extracted
        ],
    }
    state["materials"].append(material)
    save_state(state)
    return {"message": "Upload processed successfully.", "material": material}


@app.post("/api/tutor/ask")
def ask_question(payload: dict[str, Any]):
    question = payload.get("question") or ""
    user_id = payload.get("userId", 1)
    level = payload.get("level", "beginner")
    if not question.strip():
        raise HTTPException(status_code=400, detail="Please enter a question.")
    state = load_state()
    user_materials = [item for item in state["materials"] if item["uploadedByUserId"] == user_id]
    result = build_answer(question, level, user_materials)
    result["userId"] = user_id
    return result


@app.post("/api/quizzes/generate")
def generate_quiz(payload: dict[str, Any]):
    topic = payload.get("topic", "Normalization")
    difficulty = payload.get("difficulty", "medium")
    count = int(payload.get("count", 5))
    questions = compute_quiz_questions(topic, difficulty, count)
    quiz = {
        "id": f"quiz-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}",
        "topic": topic,
        "difficulty": difficulty,
        "questions": questions,
        "createdAt": utc_now(),
    }
    state = load_state()
    state.setdefault("quizBank", []).append(quiz)
    save_state(state)
    return {"quiz": quiz}


@app.post("/api/quizzes/submit")
def submit_quiz(payload: dict[str, Any]):
    quiz_id = payload.get("quizId")
    user_id = payload.get("userId", 1)
    responses = payload.get("responses", {})
    state = load_state()
    quiz = next((item for item in state.get("quizBank", []) if item["id"] == quiz_id), None)
    if not quiz:
        raise HTTPException(status_code=404, detail="Quiz not found.")

    total = 0
    correct = 0
    detailed = []
    for question in quiz["questions"]:
        total += 1
        answer = responses.get(str(question["question"]))
        is_correct = answer == question["answer"]
        if is_correct:
            correct += 1
        detailed.append({
            "question": question["question"],
            "yourAnswer": answer,
            "correct": is_correct,
            "explanation": question["explanation"],
        })

    accuracy = (correct / total * 100) if total else 0
    update_mastery(user_id, quiz["topic"], accuracy / 100)
    weak_topic = quiz["topic"]
    misconceptions = []
    if accuracy < 75:
        misconceptions.append({
            "topic": weak_topic,
            "likelyMisconception": "You may be treating normalization as a formatting rule instead of a way to remove duplicate data and prevent update anomalies.",
            "example": "A StudentCourse table storing student names in every row is redundant; splitting the data into separate tables fixes this.",
            "followUpQuestion": "If a student changes their name, which table should be updated once rather than repeated across many rows?",
            "status": "possible misconception",
        })

    score_record = {"quizId": quiz_id, "userId": user_id, "topic": quiz["topic"], "score": round(accuracy), "submittedAt": utc_now()}
    state["quizHistory"].append(score_record)
    save_state(state)
    return {"score": round(accuracy), "correct": correct, "total": total, "details": detailed, "misconceptions": misconceptions}


@app.get("/api/recommendations/{user_id}")
def recommendations(user_id: int):
    return build_recommendations(user_id)


@app.get("/api/revision/{user_id}")
def revision_tasks(user_id: int):
    state = load_state()
    return [task for task in state["revisionTasks"] if task["userId"] == user_id]


@app.post("/api/revision")
def create_revision_task(payload: dict[str, Any]):
    state = load_state()
    task = {
        "id": f"task-{len(state['revisionTasks']) + 1}",
        "userId": payload.get("userId", 1),
        "title": payload.get("title", "New revision task"),
        "topic": payload.get("topic", "General"),
        "dueDate": payload.get("dueDate", "2026-10-18"),
        "status": payload.get("status", "pending"),
    }
    state["revisionTasks"].append(task)
    save_state(state)
    return {"task": task}


@app.patch("/api/revision/{task_id}")
def update_revision_task(task_id: str, payload: dict[str, Any]):
    state = load_state()
    task = next((item for item in state["revisionTasks"] if item["id"] == task_id), None)
    if not task:
        raise HTTPException(status_code=404, detail="Task not found.")
    allowed_fields = {"title", "topic", "dueDate", "status"}
    task.update({key: value for key, value in payload.items() if key in allowed_fields})
    save_state(state)
    return {"task": task}


@app.get("/api/analytics/{user_id}")
def analytics(user_id: int):
    return build_analytics(user_id)


@app.get("/api/settings")
def settings():
    return {"theme": "default", "language": "English", "notifications": True, "demoMode": True}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
